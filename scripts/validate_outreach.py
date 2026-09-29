#!/usr/bin/env python3
"""Validate a generic professor-outreach artifact without embedding applicant data."""

import argparse
import copy
import json
import re
import sys
from pathlib import Path


PLACEHOLDER_RE = re.compile(r"\{\{[^{}]+\}\}|\[[^\[\]\n]{2,80}\]")
EMAIL_RE = re.compile(r"(?i)\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b")
CN_MOBILE_RE = re.compile(r"(?<!\d)1[3-9]\d{9}(?!\d)")
CN_ID_RE = re.compile(r"(?<!\d)\d{17}[\dXx](?!\d)")
URL_RE = re.compile(r"(?i)https?://|www\.")
ZERO_WIDTH_RE = re.compile("[\u200b\u200c\u200d\ufeff]")
FIT_LABELS = {"direct", "adjacent", "exploratory"}


def fail(errors):
    for error in errors:
        print(f"ERROR: {error}")
    return 1


def paragraph_text(item):
    if isinstance(item, str):
        return item
    if not isinstance(item, dict):
        return ""
    if isinstance(item.get("text"), str):
        return item["text"]
    runs = item.get("runs", [])
    if not isinstance(runs, list):
        return ""
    return "".join(
        run.get("text", "") for run in runs if isinstance(run, dict)
    )


def load_artifact(path):
    suffix = path.suffix.lower()
    metadata = {}
    if suffix == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data, dict):
            raise ValueError("artifact JSON root must be an object")
        metadata = data.get("metadata", {})
        if not isinstance(metadata, dict):
            raise ValueError("artifact metadata must be an object")
        items = data.get("paragraphs")
        if not isinstance(items, list) or not items:
            raise ValueError("artifact JSON must contain a non-empty paragraphs array")
        paragraphs = [paragraph_text(item).strip() for item in items]
    elif suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError("python-docx is required to validate DOCX files") from exc
        document = Document(path)
        paragraphs = [item.text.strip() for item in document.paragraphs]
    elif suffix in {".md", ".txt"}:
        text = path.read_text(encoding="utf-8")
        paragraphs = [item.strip() for item in re.split(r"\n\s*\n", text)]
    else:
        raise ValueError("artifact must be .json, .docx, .md, or .txt")
    return [item for item in paragraphs if item], metadata


def load_profile(path):
    data = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(data, dict):
        raise ValueError("profile JSON root must be an object")
    return data


def string_values(value):
    if isinstance(value, str):
        yield value
    elif isinstance(value, dict):
        for item in value.values():
            yield from string_values(item)
    elif isinstance(value, list):
        for item in value:
            yield from string_values(item)


def validate_profile(profile):
    errors = []
    for field in ("name", "affiliation", "current_status"):
        if not isinstance(profile.get(field), str) or not profile[field].strip():
            errors.append(f"profile.{field} must be a non-empty string")

    target = profile.get("target")
    if not isinstance(target, dict):
        errors.append("profile.target must be an object")
    else:
        for field in ("cycle", "role"):
            if not isinstance(target.get(field), str) or not target[field].strip():
                errors.append(f"profile.target.{field} must be a non-empty string")

    proof_points = profile.get("proof_points")
    if not isinstance(proof_points, list) or not proof_points:
        errors.append("profile.proof_points must contain at least one verified item")
    else:
        for index, item in enumerate(proof_points, start=1):
            if not isinstance(item, dict):
                errors.append(f"profile.proof_points[{index}] must be an object")
                continue
            for field in ("topic", "role", "status", "source"):
                if not isinstance(item.get(field), str) or not item[field].strip():
                    errors.append(
                        f"profile.proof_points[{index}].{field} must be a non-empty string"
                    )

    attachments = profile.get("attachments_confirmed", [])
    if not isinstance(attachments, list) or not all(
        isinstance(item, str) and item.strip() for item in attachments
    ):
        errors.append("profile.attachments_confirmed must be an array of non-empty strings")

    contact = profile.get("contact", {})
    if not isinstance(contact, dict):
        errors.append("profile.contact must be an object when provided")

    if any(PLACEHOLDER_RE.search(value) for value in string_values(profile)):
        errors.append("profile contains unresolved bracket or mustache placeholders")
    return errors


def validate_artifact(paragraphs, metadata, profile=None, min_length=None, max_length=None):
    errors = []
    text = "\n".join(paragraphs)

    if PLACEHOLDER_RE.search(text):
        errors.append("email contains unresolved bracket or mustache placeholders")
    if ZERO_WIDTH_RE.search(text):
        errors.append("email contains zero-width characters")
    if URL_RE.search(text):
        errors.append("email body contains a URL; keep evidence links outside the body")
    if CN_ID_RE.search(text):
        errors.append("email contains a possible identity-card number")
    if re.search(r"(?:^|\s)\[\d+\](?:\s|$)", text):
        errors.append("email contains a citation marker")

    titles = re.findall(r"《([^》]+)》", text)
    if len(titles) > 1:
        errors.append(f"email names {len(titles)} hook papers; expected at most one")

    professor_name = metadata.get("professor_name")
    if professor_name is not None:
        if not isinstance(professor_name, str) or not professor_name.strip():
            errors.append("metadata.professor_name must be a non-empty string")
        elif not paragraphs or professor_name not in paragraphs[0]:
            errors.append("professor name is missing from the salutation")

    hook_paper = metadata.get("hook_paper")
    if hook_paper is not None:
        if not isinstance(hook_paper, str) or not hook_paper.strip():
            errors.append("metadata.hook_paper must be a non-empty string")
        elif hook_paper not in titles:
            errors.append("metadata.hook_paper does not match the paper title in the email")

    fit_label = metadata.get("fit_label")
    if fit_label is not None and fit_label not in FIT_LABELS:
        errors.append("metadata.fit_label must be direct, adjacent, or exploratory")

    mentioned = metadata.get("attachments_mentioned", [])
    if not isinstance(mentioned, list) or not all(
        isinstance(item, str) and item.strip() for item in mentioned
    ):
        errors.append("metadata.attachments_mentioned must be an array of non-empty strings")
        mentioned = []

    if profile is not None:
        errors.extend(validate_profile(profile))
        for field in ("name", "affiliation"):
            value = profile.get(field)
            if isinstance(value, str) and value.strip() and value not in text:
                errors.append(f"profile.{field} is missing from the email")

        target = profile.get("target", {})
        if isinstance(target, dict):
            for field in ("cycle", "role"):
                value = target.get(field)
                if isinstance(value, str) and value.strip() and value not in text:
                    errors.append(f"profile.target.{field} is missing from the email")

        confirmed = profile.get("attachments_confirmed", [])
        if isinstance(confirmed, list):
            unknown = [item for item in mentioned if item not in confirmed]
            if unknown:
                errors.append("email mentions attachments not confirmed by the profile")

        contact = profile.get("contact", {})
        if isinstance(contact, dict):
            allowed_emails = {
                value.casefold()
                for key, value in contact.items()
                if "email" in key.casefold() and isinstance(value, str) and value.strip()
            }
            found_emails = {value.casefold() for value in EMAIL_RE.findall(text)}
            if not found_emails.issubset(allowed_emails):
                errors.append("email body contains an email address not present in profile.contact")

            allowed_phones = {
                match.group(0)
                for value in contact.values()
                if isinstance(value, str)
                for match in CN_MOBILE_RE.finditer(value)
            }
            found_phones = set(CN_MOBILE_RE.findall(text))
            if not found_phones.issubset(allowed_phones):
                errors.append("email body contains a phone number not present in profile.contact")

        if not mentioned and re.search(r"附件|随信附上|attach(?:ed|ment)", text, re.I):
            confirmed = profile.get("attachments_confirmed", [])
            if not confirmed:
                errors.append("email refers to an attachment, but the profile confirms none")

    language = str(metadata.get("language", "")).casefold()
    is_chinese = language in {"chinese", "zh", "zh-cn"} or (
        not language and len(re.findall(r"[\u4e00-\u9fff]", text)) >= 20
    )
    measured_length = (
        len(re.findall(r"[\u4e00-\u9fffA-Za-z0-9]", text))
        if is_chinese
        else len(re.findall(r"\b[\w'-]+\b", text))
    )
    lower = min_length if min_length is not None else (250 if is_chinese else 120)
    upper = max_length if max_length is not None else (900 if is_chinese else 420)
    unit = "characters" if is_chinese else "words"
    if measured_length < lower:
        errors.append(f"email is too short: {measured_length} {unit}; minimum is {lower}")
    if measured_length > upper:
        errors.append(f"email is too long: {measured_length} {unit}; maximum is {upper}")

    return errors


def run_self_test():
    profile = {
        "name": "申请人甲",
        "affiliation": "示例大学示例学院",
        "current_status": "本科生",
        "target": {"cycle": "2027", "role": "博士生", "language": "Chinese"},
        "proof_points": [
            {
                "topic": "示例研究",
                "role": "主要参与者",
                "methods": ["数据分析"],
                "outcome": "完成验证",
                "status": "completed",
                "source": "user-supplied CV",
            }
        ],
        "attachments_confirmed": ["个人简历"],
        "contact": {},
    }
    paragraphs = [
        "尊敬的示例教授老师：",
        "您好！我是示例大学示例学院本科生申请人甲，计划申请2027博士生项目。",
        "我在示例研究中担任主要参与者，围绕数据分析完成了方案设计与实验验证，并形成了可复核的结果。相关经历使我能够从问题定义、数据处理到结果解释完整推进研究任务。",
        "我了解到您持续研究可信机器学习。您团队的《示例论文》尤其吸引我，该工作通过稳健表征学习改善复杂分布下的模型可靠性。我的示例研究与其应用场景不同，但数据分析和实验验证经验为进一步研究模型可靠性评估提供了基础。",
        "我的个人简历已随信附上。若您认为我的背景与团队方向存在可讨论的结合点，我希望有机会进一步汇报，并请教博士生申请与后续研究安排。",
        "祝工作顺利！申请人甲，示例大学示例学院",
    ]
    metadata = {
        "professor_name": "示例教授",
        "hook_paper": "示例论文",
        "fit_label": "adjacent",
        "language": "Chinese",
        "attachments_mentioned": ["个人简历"],
    }
    valid_errors = validate_artifact(
        paragraphs, metadata, profile=profile, min_length=100, max_length=1000
    )
    if valid_errors:
        return fail([f"self-test valid case failed: {item}" for item in valid_errors])

    invalid_paragraphs = copy.deepcopy(paragraphs)
    invalid_paragraphs[1] += " [未填写字段]"
    invalid_errors = validate_artifact(
        invalid_paragraphs, metadata, profile=profile, min_length=100, max_length=1000
    )
    if not any("placeholder" in item for item in invalid_errors):
        return fail(["self-test invalid case was not rejected"])
    print("Generic outreach validator self-test passed.")
    return 0


def main():
    parser = argparse.ArgumentParser(
        description="Validate a generic professor-outreach email without a bundled personal preset.",
        epilog=(
            "JSON artifacts use {'metadata': {'professor_name', 'hook_paper', "
            "'fit_label', 'language', 'attachments_mentioned'}, 'paragraphs': "
            "[string or {'text': string} or {'runs': [{'text': string}]}]}. "
            "Profile JSON follows references/applicant_profile.md."
        ),
    )
    parser.add_argument("artifact", nargs="?", type=Path)
    parser.add_argument("--profile", type=Path, help="Optional private .local.json profile")
    parser.add_argument("--min-length", type=int)
    parser.add_argument("--max-length", type=int)
    parser.add_argument("--self-test", action="store_true")
    args = parser.parse_args()

    if args.self_test:
        return run_self_test()
    if args.artifact is None:
        parser.error("artifact is required unless --self-test is used")
    if args.min_length is not None and args.min_length < 0:
        parser.error("--min-length must be non-negative")
    if args.max_length is not None and args.max_length < 1:
        parser.error("--max-length must be positive")
    if (
        args.min_length is not None
        and args.max_length is not None
        and args.min_length > args.max_length
    ):
        parser.error("--min-length cannot exceed --max-length")

    try:
        paragraphs, metadata = load_artifact(args.artifact)
        profile = load_profile(args.profile) if args.profile else None
    except (OSError, ValueError, json.JSONDecodeError) as exc:
        return fail([str(exc)])

    errors = validate_artifact(
        paragraphs,
        metadata,
        profile=profile,
        min_length=args.min_length,
        max_length=args.max_length,
    )
    if errors:
        return fail(errors)
    print("Generic outreach validation passed.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
